"""Helpers and approved categories for the Yemeni laws archive."""

import io
from typing import Optional


LEGAL_CATEGORIES = {
    "constitution": "الدستور اليمني",
    "civil_law": "القانون المدني",
    "criminal_law": "قانون الجرائم والعقوبات",
    "procedures_law": "قانون المرافعات والتنفيذ",
    "commercial_law": "القانون التجاري",
    "companies_law": "قانون الشركات",
    "labor_law": "قانون العمل",
    "family_law": "قانون الأحوال الشخصية",
    "investment_law": "قوانين الاستثمار والضرائب",
    "real_estate_law": "قوانين الأراضي والعقارات",
    "regulations": "اللوائح والقرارات التنفيذية",
    "other_yemeni_law": "قوانين ولوائح يمنية أخرى",
}

ALLOWED_LAW_TYPES = {"law", "regulation"}
ALLOWED_FILE_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024


def extract_legal_text(data: bytes, extension: str) -> str:
    """Extract searchable text from common legal document formats, without AI."""
    extension = extension.lower()
    if extension == ".txt":
        return data.decode("utf-8-sig", errors="replace").strip()
    if extension == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs).strip()
    return ""


def law_excerpt(text: Optional[str], query: str, limit: int = 360) -> str:
    text = (text or "").strip()
    if not text:
        return ""
    lowered = text.casefold()
    position = lowered.find(query.casefold()) if query else -1
    start = max(0, position - limit // 3) if position >= 0 else 0
    excerpt = text[start:start + limit]
    return ("…" if start else "") + excerpt + ("…" if start + limit < len(text) else "")
